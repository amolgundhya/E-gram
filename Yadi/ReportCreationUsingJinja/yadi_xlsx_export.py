"""नमुना-8 रजिस्टर / All Report - खरा .xlsx export (CSV ऐवजी).

Sheet 1 "मालमत्ता": एक ओळ प्रति मालमत्ता - बोल्ड हेडर, AutoFilter, गोठलेली वरची ओळ,
आकडे खरे नंबर म्हणून (2 दशांश), तळाशी बेरजेची ओळ.
Sheet 2 "बांधकाम तपशील": एक ओळ प्रति बांधकाम रो (भारांक/वापर सकट), मालमत्ता क्रमांकाने जोडलेली.

पासवर्ड सेट असेल तरच msoffcrypto-tool ने .xlsx एन्क्रिप्ट करतो; नसेल तर साधा .xlsx.
"""
import io
from openpyxl import Workbook
from openpyxl.styles import Font, Alignment
from openpyxl.utils import get_column_letter

SHEET1_COLUMNS = [
    ("srNo", "अ.क्र.", "int"),
    ("propertyNumber", "मालमत्ता क्रमांक", "text"),
    ("ownerName", "मालमत्ता धारकाचे नाव", "text"),
    ("ownerWifeName", "पत्नीचे नाव", "text"),
    ("occupantName", "भोगवटदार", "text"),
    ("aadharNumber", "आधार क्रमांक", "text"),
    ("mobileNumber", "मोबाईल क्रमांक", "text"),
    ("village", "गाव", "text"),
    ("roadName", "रस्त्याचे नाव", "text"),
    ("cityWardGatNumber", "सिटी सर्वे / गट नं", "text"),
    ("areaEast", "पूर्व लांबी", "num"),
    ("areaWest", "पश्चिम लांबी", "num"),
    ("areaNorth", "उत्तर लांबी", "num"),
    ("areaSouth", "दक्षिण लांबी", "num"),
    ("boundaryEast", "पूर्व सीमा", "text"),
    ("boundaryWest", "पश्चिम सीमा", "text"),
    ("boundaryNorth", "उत्तर सीमा", "text"),
    ("boundarySouth", "दक्षिण सीमा", "text"),
    ("dwarNaav", "व्दार", "text"),
    ("housingUnit", "एकक", "text"),
    ("total_arearinfoot", "एकूण क्षेत्रफळ (चौ.फु.)", "num"),
    ("totalareainmeters", "एकूण क्षेत्रफळ (चौ.मी.)", "num"),
    ("constructionCount", "बांधकाम संख्या", "int"),
    ("totalconstructionareainfoot", "एकूण बांधकाम क्षेत्रफळ (चौ.फु.)", "num"),
    ("khaliJagaFt", "खाली जागा (चौ.फु.)", "num"),
    ("khaliJagaM", "खाली जागा (चौ.मी.)", "num"),
    ("totalCapitalValue", "भांडवली मूल्य", "num"),
    ("totalHouseTax", "घरपट्टी", "num"),
    ("lightingTax", "दिवाबत्ती कर", "num"),
    ("healthTax", "आरोग्य कर", "num"),
    ("sapanikar", "सा.पाणी कर", "num"),
    ("vpanikar", "वि.पाणी कर", "num"),
    ("cleaningTax", "सफाई कर", "num"),
    ("toiletTax", "शौचालय कर", "num"),
    ("totaltax", "एकूण कर", "num"),
    ("removeLightHealthTax", "दिवा/आरोग्य कर माफ", "bool_yn"),
    ("applyCleaningTax", "सफाई कर लागू", "bool_yn"),
    ("applyToiletTax", "शौचालय कर लागू", "bool_yn"),
    ("taxNotApplicable", "कर लागू नाही", "bool_yn"),
    ("toilet", "शौचालय", "text"),
    ("toiletBenefitYear", "शौचालय लाभ वर्ष", "text"),
    ("gharkul", "घरकुल", "text"),
    ("gharkulYojana", "घरकुल योजना", "text"),
    ("gharkulBenefitYear", "घरकुल लाभ वर्ष", "text"),
    ("waterFacility1", "पाणी सुविधा १", "text"),
    ("waterFacility2", "पाणी सुविधा २", "text"),
    ("house", "छप्पर प्रकार", "text"),
    ("remarks", "शेरा", "text"),
]

SHEET1_TOTAL_COLUMNS = {
    "total_arearinfoot", "totalareainmeters", "totalconstructionareainfoot",
    "khaliJagaFt", "khaliJagaM", "totalCapitalValue", "totalHouseTax",
    "lightingTax", "healthTax", "sapanikar", "vpanikar", "cleaningTax",
    "toiletTax", "totaltax",
}

# Aadhar/mobile numbers must never be written as Excel numeric cells - a long
# all-digit number in a "General" numeric cell gets auto-displayed in
# scientific notation (e.g. 9.11E+11). Forcing these to text cells (both the
# Python value and the cell's number format) keeps them exactly as typed.
FORCE_TEXT_COLUMNS = {"aadharNumber", "mobileNumber"}

SHEET2_COLUMNS = [
    ("srNo", "अ.क्र.", "int"),
    ("propertyNumber", "मालमत्ता क्रमांक", "text"),
    ("type", "बांधकाम प्रकार", "text"),
    ("length", "लांबी", "num"),
    ("width", "रुंदी", "num"),
    ("unit", "एकक", "text"),
    ("area", "क्षेत्रफळ", "num"),
    ("year", "बांधकाम वर्ष", "text"),
    ("floor", "मजला", "text"),
    ("weightage", "भारांक", "num"),
    ("usage", "वापर", "text"),
]

HEADER_FONT = Font(bold=True)
HEADER_ALIGN = Alignment(horizontal="center", vertical="center", wrap_text=True)
TOTAL_FONT = Font(bold=True)


def _num(value):
    try:
        if value is None or value == "":
            return None
        return round(float(value), 2)
    except (TypeError, ValueError):
        return None


EX_SERVICEMAN_NOTE = (
    "आजी / माजी सैनिकांना ग्रामपंचायत मालमत्ता करातून सुट देण्याबाबत शासन परिपत्रक "
    "व्हीपीएम-2020/प्र.क्र.93/पंरा-4, दि. 18 ऑगस्ट 2020 नुसार."
)


def _remarks_with_ex_serviceman_note(prop):
    """शेरा (remarks) कॉलम - मालमत्ता आजी/माजी सैनिक टिक असेल तर टीप जोडते, जुना मजकूर कायम ठेवून."""
    remarks = prop.get("remarks")
    if not prop.get("exServicemanTip"):
        return remarks
    return f"{remarks} | {EX_SERVICEMAN_NOTE}" if remarks else EX_SERVICEMAN_NOTE


def _setup_header(ws, columns):
    for col_idx, (_, label, _) in enumerate(columns, start=1):
        cell = ws.cell(row=1, column=col_idx, value=label)
        cell.font = HEADER_FONT
        cell.alignment = HEADER_ALIGN
    last_col_letter = get_column_letter(len(columns))
    ws.auto_filter.ref = f"A1:{last_col_letter}1"
    ws.freeze_panes = "A2"


def _autosize_columns(ws, columns, sample_rows=200):
    for col_idx, (_, label, _) in enumerate(columns, start=1):
        letter = get_column_letter(col_idx)
        max_len = len(str(label))
        for row in ws.iter_rows(min_col=col_idx, max_col=col_idx, min_row=2, max_row=min(ws.max_row, sample_rows + 1)):
            v = row[0].value
            if v is not None:
                max_len = max(max_len, len(str(v)))
        ws.column_dimensions[letter].width = min(max(max_len + 2, 10), 40)


def build_namuna8_register_workbook(properties: list) -> Workbook:
    wb = Workbook()

    # ---------------- Sheet 1: मालमत्ता ----------------
    ws1 = wb.active
    ws1.title = "मालमत्ता"
    _setup_header(ws1, SHEET1_COLUMNS)

    totals = {key: 0.0 for key in SHEET1_TOTAL_COLUMNS}
    row_idx = 2
    for prop in properties:
        construction_list = prop.get("constructionType") or []
        khali_list = prop.get("khaliJaga") or []
        khali_ft = khali_list[0].get("totalkhalijagaareainfoot") if khali_list else None
        khali_m = khali_list[0].get("totalkhalijagaareainmeters") if khali_list else None

        row_values = {
            "srNo": prop.get("srNo"),
            "propertyNumber": prop.get("propertyNumber"),
            "ownerName": prop.get("ownerName"),
            "ownerWifeName": prop.get("ownerWifeName"),
            "occupantName": prop.get("occupantName"),
            "aadharNumber": prop.get("aadharNumber"),
            "mobileNumber": prop.get("mobileNumber"),
            "village": prop.get("village"),
            "roadName": prop.get("roadName"),
            "cityWardGatNumber": prop.get("cityWardGatNumber"),
            "areaEast": prop.get("areaEast"),
            "areaWest": prop.get("areaWest"),
            "areaNorth": prop.get("areaNorth"),
            "areaSouth": prop.get("areaSouth"),
            "boundaryEast": prop.get("boundaryEast"),
            "boundaryWest": prop.get("boundaryWest"),
            "boundaryNorth": prop.get("boundaryNorth"),
            "boundarySouth": prop.get("boundarySouth"),
            "dwarNaav": prop.get("dwarNaav"),
            "housingUnit": prop.get("housingUnit") or prop.get("areaUnit"),
            "total_arearinfoot": prop.get("total_arearinfoot"),
            "totalareainmeters": prop.get("totalareainmeters"),
            "constructionCount": len(construction_list),
            "totalconstructionareainfoot": prop.get("totalconstructionareainfoot"),
            "khaliJagaFt": khali_ft,
            "khaliJagaM": khali_m,
            "totalCapitalValue": prop.get("totalCapitalValue"),
            "totalHouseTax": prop.get("totalHouseTax"),
            "lightingTax": prop.get("lightingTax"),
            "healthTax": prop.get("healthTax"),
            "sapanikar": prop.get("sapanikar"),
            "vpanikar": prop.get("vpanikar"),
            "cleaningTax": prop.get("cleaningTax"),
            "toiletTax": prop.get("toiletTax"),
            "totaltax": prop.get("totaltax"),
            "removeLightHealthTax": prop.get("removeLightHealthTax"),
            "applyCleaningTax": prop.get("applyCleaningTax"),
            "applyToiletTax": prop.get("applyToiletTax"),
            "taxNotApplicable": prop.get("taxNotApplicable"),
            "toilet": prop.get("toilet"),
            "toiletBenefitYear": prop.get("toiletBenefitYear"),
            "gharkul": prop.get("gharkul"),
            "gharkulYojana": prop.get("gharkulYojana"),
            "gharkulBenefitYear": prop.get("gharkulBenefitYear"),
            "waterFacility1": prop.get("waterFacility1"),
            "waterFacility2": prop.get("waterFacility2"),
            "house": prop.get("house"),
            "remarks": _remarks_with_ex_serviceman_note(prop),
        }

        for col_idx, (key, _, kind) in enumerate(SHEET1_COLUMNS, start=1):
            value = row_values.get(key)
            if kind == "num":
                value = _num(value)
                if key in SHEET1_TOTAL_COLUMNS and value is not None:
                    totals[key] += value
                cell = ws1.cell(row=row_idx, column=col_idx, value=value)
                if value is not None:
                    cell.number_format = "0.00"
            elif kind == "int":
                cell = ws1.cell(row=row_idx, column=col_idx, value=value if value not in (None, "") else None)
            elif kind == "bool_yn":
                cell = ws1.cell(row=row_idx, column=col_idx, value="होय" if value else "नाही")
            else:
                text_value = str(value) if key in FORCE_TEXT_COLUMNS and value not in (None, "") else (value if value not in (None, "") else None)
                cell = ws1.cell(row=row_idx, column=col_idx, value=text_value)
                if key in FORCE_TEXT_COLUMNS:
                    cell.number_format = "@"
        row_idx += 1

    # तळाशी बेरजेची ओळ
    total_row = row_idx
    label_cell = ws1.cell(row=total_row, column=1, value="एकूण")
    label_cell.font = TOTAL_FONT
    ws1.merge_cells(start_row=total_row, start_column=1, end_row=total_row, end_column=8)
    for col_idx, (key, _, kind) in enumerate(SHEET1_COLUMNS, start=1):
        if key in SHEET1_TOTAL_COLUMNS:
            cell = ws1.cell(row=total_row, column=col_idx, value=round(totals[key], 2))
            cell.font = TOTAL_FONT
            cell.number_format = "0.00"
    # बेरीज ओळ फिल्टर/स्क्रोलमुळे चुकून डेटा समजू नये म्हणून थोडी वेगळी दिसावी.
    for col_idx in range(1, len(SHEET1_COLUMNS) + 1):
        ws1.cell(row=total_row, column=col_idx).font = TOTAL_FONT

    _autosize_columns(ws1, SHEET1_COLUMNS)

    # ---------------- Sheet 2: बांधकाम तपशील ----------------
    ws2 = wb.create_sheet("बांधकाम तपशील")
    _setup_header(ws2, SHEET2_COLUMNS)

    row_idx = 2
    sr = 1
    for prop in properties:
        construction_list = prop.get("constructionType") or []
        construction_unit = prop.get("constructionAreaUnit") or prop.get("housingUnit") or prop.get("areaUnit")
        for item in construction_list:
            length = _num(item.get("length"))
            width = _num(item.get("width"))
            area = round(length * width, 2) if (length is not None and width is not None) else None
            row_values = {
                "srNo": sr,
                "propertyNumber": prop.get("propertyNumber"),
                "type": item.get("type"),
                "length": length,
                "width": width,
                "unit": construction_unit,
                "area": area,
                "year": item.get("year"),
                "floor": item.get("floor"),
                "weightage": _num(item.get("usageBasedBuildingWeightageFactor")),
                "usage": item.get("usage"),
            }
            for col_idx, (key, _, kind) in enumerate(SHEET2_COLUMNS, start=1):
                value = row_values.get(key)
                cell = ws2.cell(row=row_idx, column=col_idx, value=value if value not in (None, "") else None)
                if kind == "num" and value is not None:
                    cell.number_format = "0.00"
            row_idx += 1
            sr += 1

    _autosize_columns(ws2, SHEET2_COLUMNS)

    return wb


def workbook_to_bytes(wb: Workbook, password: str = None) -> bytes:
    buf = io.BytesIO()
    wb.save(buf)
    buf.seek(0)
    if not password:
        return buf.getvalue()

    import msoffcrypto
    office_file = msoffcrypto.OfficeFile(buf)
    encrypted = io.BytesIO()
    office_file.encrypt(password, encrypted)
    return encrypted.getvalue()
