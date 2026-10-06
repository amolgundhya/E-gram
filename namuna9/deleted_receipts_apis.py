"""रद्द / डिलीट केलेल्या पावत्या - रिपोर्ट.

नमुना-9 (namuna9_receipts) आणि नमुना-7 (namuna7) या दोन्हीमधून is_deleted=True
असलेल्या पावत्या एकत्र आणून दाखवतो - गाव/साल/दिनांक-रेंज/प्रकार या फिल्टरसह.
फक्त वाचण्यासाठी - इथे कुठलीही पावती परत आणली (restore) जात नाही.
"""
import io
from datetime import datetime, timedelta
from typing import Optional

from fastapi import APIRouter, Depends, Query
from fastapi.responses import StreamingResponse
from sqlalchemy.orm import Session
from openpyxl import Workbook
from openpyxl.styles import Font, Alignment
from openpyxl.utils import get_column_letter

import database
from namuna9 import namuna9_model
from namuna8 import namuna8_model
from namuna8.namuna7.namuna7_model import Namuna7
from namuna8.recordresponses.property_record_response import compute_display_sr_no
from Yadi.ReportCreationUsingJinja.yadi_xlsx_export import workbook_to_bytes

router = APIRouter(prefix="/reports/deleted-receipts", tags=["Deleted Receipts Report"])

COLUMNS = [
    ("type", "प्रकार"),
    ("pavtiKramank", "पावती क्र."),
    ("pavtiBookKramank", "पावती बुक क्र."),
    ("pavtiDate", "पावती दिनांक"),
    ("srNo", "अ.क्र."),
    ("malmattaKramank", "मालमत्ता क्रमांक"),
    ("ownerName", "मालक नाव"),
    ("ghar", "घर कर"),
    ("diva", "दिवाबत्ती कर"),
    ("aarogyaKar", "आरोग्य कर"),
    ("sapanikar", "सा.पाणी कर"),
    ("vpanikar", "वि.पाणी कर"),
    ("cleaningTax", "सफाई कर"),
    ("dand", "दंड"),
    ("noticeFee", "नोटीस फी"),
    ("warrantFee", "वारंट फी"),
    ("total", "एकूण"),
    ("deletedAt", "डिलीट दिनांक-वेळ"),
    ("deletedBy", "डिलीट करणारा"),
    ("deleteReason", "कारण"),
]


def _parse_date(s: Optional[str]):
    if not s:
        return None
    try:
        return datetime.strptime(s, "%Y-%m-%d")
    except Exception:
        return None


def _fetch_rows(
    db: Session,
    district_id: Optional[int],
    taluka_id: Optional[int],
    gram_panchayat_id: Optional[int],
    village_id: Optional[int],
    year: Optional[str],
    from_date: Optional[str],
    to_date: Optional[str],
    receipt_type: Optional[str],
):
    fd = _parse_date(from_date)
    td = _parse_date(to_date)
    if td:
        td = td + timedelta(days=1)  # to_date समावेशक (inclusive) हवा

    rows = []

    if receipt_type in (None, "", "namuna9"):
        q = db.query(namuna9_model.Namuna9Receipt).filter(
            namuna9_model.Namuna9Receipt.is_deleted == True
        )
        # टिप: Namuna9.district_id/taluka_id हे जुन्या रेकॉर्ड्समध्ये विश्वासार्हपणे
        # भरलेले नसतात (बरेच रो NULL) - त्यामुळे इतर सर्व namuna9 receipt-listing
        # एंडपॉईंट्स (list_receipts, list_receipts_by_date_village/all) प्रमाणेच
        # फक्त gram_panchayat_id + village_id नुसारच स्कोप करायचं, district/taluka
        # नुसार नाही - अन्यथा खरे रो चुकून गाळले जातात.
        if gram_panchayat_id:
            q = q.filter(namuna9_model.Namuna9Receipt.gram_panchayat_id == gram_panchayat_id)
        if village_id:
            prop_ids_subq = db.query(namuna8_model.Property.id).filter(
                namuna8_model.Property.village_id == village_id
            ).subquery()
            q = q.filter(namuna9_model.Namuna9Receipt.property_id.in_(prop_ids_subq))
        if year:
            namuna9_ids_subq = db.query(namuna9_model.Namuna9.id).filter(
                namuna9_model.Namuna9.yearslap == year
            ).subquery()
            q = q.filter(namuna9_model.Namuna9Receipt.namuna9_id.in_(namuna9_ids_subq))
        if fd:
            q = q.filter(namuna9_model.Namuna9Receipt.deleted_at >= fd)
        if td:
            q = q.filter(namuna9_model.Namuna9Receipt.deleted_at < td)

        for r in q.all():
            prop = db.query(namuna8_model.Property).filter(
                namuna8_model.Property.id == r.property_id
            ).first()
            rows.append({
                "type": "नमुना-9",
                "pavtiKramank": r.pavti_kramank,
                "pavtiBookKramank": r.pa_book_kramank,
                "pavtiDate": r.pavti_date.strftime("%Y-%m-%d") if r.pavti_date else None,
                "srNo": compute_display_sr_no(db, prop.village_id, prop.anuKramank) if prop else None,
                "malmattaKramank": r.malmatta_kramank,
                "ownerName": r.owner_name,
                "ghar": round((r.vasuliGhar or 0) + (r.vasuliChaluGhar or 0), 2),
                "diva": round((r.vasuliDiva or 0) + (r.vasuliChaluDiva or 0), 2),
                "aarogyaKar": round((r.vasuliAarogyaKar or 0) + (r.vasuliChaluAarogyaKar or 0), 2),
                "sapanikar": round((r.vasuliSapanikar or 0) + (r.vasuliChaluSapanikar or 0), 2),
                "vpanikar": round((r.vasuliVpanikar or 0) + (r.vasuliChaluVpanikar or 0), 2),
                "cleaningTax": round((r.vasuliCleaningTax or 0) + (r.vasuliChaluCleaningTax or 0), 2),
                "dand": round(r.vasuliDand or 0, 2),
                "noticeFee": round(r.vasuliNoticeFee or 0, 2),
                "warrantFee": round(r.vasuliWarrantFee or 0, 2),
                "total": round(r.total or 0, 2),
                "deletedAt": r.deleted_at.strftime("%Y-%m-%d %H:%M") if r.deleted_at else None,
                "deletedBy": r.deleted_by,
                "deleteReason": r.delete_reason,
            })

    if receipt_type in (None, "", "namuna7"):
        q = db.query(Namuna7).filter(Namuna7.is_deleted == True)
        if district_id:
            q = q.filter(Namuna7.district_id == district_id)
        if taluka_id:
            q = q.filter(Namuna7.taluka_id == taluka_id)
        if gram_panchayat_id:
            q = q.filter(Namuna7.gram_panchayat_id == gram_panchayat_id)
        if village_id:
            q = q.filter(Namuna7.villageId == village_id)
        # नमुना-7 ला "साल" (yearslap) नाही - year फिल्टर लागू असताना नमुना-7 रो
        # वर्गळल्या जात नाहीत, कारण तो संकल्पनाच त्या पावती-बुकला लागू नाही.
        if fd:
            q = q.filter(Namuna7.deleted_at >= fd)
        if td:
            q = q.filter(Namuna7.deleted_at < td)

        for r in q.all():
            owner = db.query(namuna8_model.Owner).filter(namuna8_model.Owner.id == r.userId).first()
            prop = owner.properties[0] if owner and owner.properties else None
            rows.append({
                "type": "नमुना-7",
                "pavtiKramank": r.receiptNumber,
                "pavtiBookKramank": r.receiptBookNumber,
                "pavtiDate": r.createdAt.strftime("%Y-%m-%d") if r.createdAt else None,
                "srNo": compute_display_sr_no(db, prop.village_id, prop.anuKramank) if prop else None,
                "malmattaKramank": (prop.malmattaKramank if prop else None) or r.malmattaKramank,
                "ownerName": owner.name if owner else None,
                "ghar": None,
                "diva": None,
                "aarogyaKar": None,
                "sapanikar": None,
                "vpanikar": None,
                "cleaningTax": None,
                "dand": None,
                "noticeFee": None,
                "warrantFee": None,
                "total": round(r.receivedMoney or 0, 2),
                "deletedAt": r.deleted_at.strftime("%Y-%m-%d %H:%M") if r.deleted_at else None,
                "deletedBy": r.deleted_by,
                "deleteReason": r.delete_reason,
            })

    rows.sort(key=lambda x: x.get("deletedAt") or "", reverse=True)
    return rows


@router.get("")
def list_deleted_receipts(
    district_id: Optional[int] = None,
    taluka_id: Optional[int] = None,
    gram_panchayat_id: Optional[int] = None,
    village_id: Optional[int] = None,
    year: Optional[str] = None,
    from_date: Optional[str] = None,
    to_date: Optional[str] = None,
    receipt_type: Optional[str] = Query(None, description="namuna9 | namuna7 | रिकामं असेल तर दोन्ही"),
    db: Session = Depends(database.get_db),
):
    return _fetch_rows(db, district_id, taluka_id, gram_panchayat_id, village_id, year, from_date, to_date, receipt_type)


@router.get("/export-xlsx")
def export_deleted_receipts_xlsx(
    district_id: Optional[int] = None,
    taluka_id: Optional[int] = None,
    gram_panchayat_id: Optional[int] = None,
    village_id: Optional[int] = None,
    year: Optional[str] = None,
    from_date: Optional[str] = None,
    to_date: Optional[str] = None,
    receipt_type: Optional[str] = None,
    db: Session = Depends(database.get_db),
):
    rows = _fetch_rows(db, district_id, taluka_id, gram_panchayat_id, village_id, year, from_date, to_date, receipt_type)

    # Master मध्ये सेट केलेला export पासवर्ड (गाव पंचायतीनुसार) - नसेल तर पासवर्डशिवाय export (इतर एक्सपोर्ट्स प्रमाणेच).
    password = None
    if gram_panchayat_id:
        checklist = db.query(namuna8_model.Namuna8SettingChecklist).filter(
            namuna8_model.Namuna8SettingChecklist.gram_panchayat_id == gram_panchayat_id
        ).first()
        if checklist and getattr(checklist, 'exportPassword', None):
            password = checklist.exportPassword

    wb = Workbook()
    ws = wb.active
    ws.title = "डिलीट पावत्या"
    header_font = Font(bold=True)
    header_align = Alignment(horizontal="center", vertical="center", wrap_text=True)
    for col_idx, (_, label) in enumerate(COLUMNS, start=1):
        cell = ws.cell(row=1, column=col_idx, value=label)
        cell.font = header_font
        cell.alignment = header_align
    ws.auto_filter.ref = f"A1:{get_column_letter(len(COLUMNS))}1"
    ws.freeze_panes = "A2"

    for row_idx, row in enumerate(rows, start=2):
        for col_idx, (key, _) in enumerate(COLUMNS, start=1):
            ws.cell(row=row_idx, column=col_idx, value=row.get(key))

    for col_idx, (_, label) in enumerate(COLUMNS, start=1):
        letter = get_column_letter(col_idx)
        ws.column_dimensions[letter].width = max(len(str(label)) + 2, 12)

    xlsx_bytes = workbook_to_bytes(wb, password=password)
    return StreamingResponse(
        io.BytesIO(xlsx_bytes),
        media_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
        headers={"Content-Disposition": 'attachment; filename="deleted_receipts.xlsx"'},
    )
